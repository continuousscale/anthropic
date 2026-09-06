"""Turn a file's bytes into content blocks the classifier can read.

PDFs and images are sent to Claude natively rather than being run through a
local OCR step. That is deliberate: the files this agent exists for are phone
photos of receipts and flatbed scans of signed agreements, and a text-only
pipeline reads those as empty. Vision sees the letterhead, the signature and
the handwritten date.

Office formats have no visual layer worth sending, so those are extracted to
text locally.
"""

from __future__ import annotations

import base64
import csv
import io
import logging
from pathlib import Path
from typing import Any

log = logging.getLogger(__name__)

# pypdf logs malformed-file complaints at error level; we handle those
# ourselves and report them per file rather than as loose console noise.
logging.getLogger("pypdf").setLevel(logging.CRITICAL)

# Only the first pages matter for identification, and they keep requests small.
MAX_PDF_PAGES = 12
# The API caps a request at 32MB; stay well under it after base64 expansion.
MAX_DOCUMENT_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 24_000

PDF_MIMES = {"application/pdf"}
IMAGE_MIMES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
TEXT_MIMES = {"text/plain", "text/markdown", "application/json", "text/html"}
CSV_MIMES = {"text/csv", "text/tab-separated-values"}
DOCX_MIMES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
XLSX_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}


class Unreadable(Exception):
    """The file cannot be turned into anything the classifier can read."""


def _trim_pdf(data: bytes) -> bytes:
    """Keep the first pages so a 200-page export still classifies cheaply."""
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError:  # pragma: no cover - dependency is declared
        return data

    try:
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) <= MAX_PDF_PAGES:
            return data
        writer = PdfWriter()
        for page in reader.pages[:MAX_PDF_PAGES]:
            writer.add_page(page)
        buf = io.BytesIO()
        writer.write(buf)
        return buf.getvalue()
    except Exception as exc:
        log.warning("Could not trim PDF (%s); sending as-is.", exc)
        return data


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in wb.worksheets:
        chunks.append(f"--- sheet: {sheet.title} ---")
        for row in sheet.iter_rows(max_row=60, values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _csv_text(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [" | ".join(r) for i, r in enumerate(reader) if i < 60]
    return "\n".join(rows)


def build_content_blocks(
    filename: str, mime_type: str, data: bytes
) -> tuple[list[dict[str, Any]], str]:
    """Build the user-content blocks describing one file.

    Returns (blocks, modality) where modality is a short label used for logging
    and for telling a reviewer how the file was read.
    """
    if not data:
        raise Unreadable("file is empty")

    mime = (mime_type or "").split(";")[0].strip().lower()
    suffix = Path(filename).suffix.lower()

    if mime in PDF_MIMES or suffix == ".pdf":
        # A file can carry a .pdf extension and not be a PDF - a renamed scan,
        # a failed export, an HTML error page saved by a browser. Sending those
        # bytes to the API wastes a call and returns nonsense.
        if not data.lstrip()[:5].startswith(b"%PDF-"):
            raise Unreadable("file has a .pdf extension but is not a PDF")
        payload = _trim_pdf(data)
        if len(payload) > MAX_DOCUMENT_BYTES:
            raise Unreadable(
                f"PDF is {len(payload) // 1024 // 1024}MB after trimming to "
                f"{MAX_PDF_PAGES} pages; too large to classify."
            )
        return (
            [
                {
                    "type": "document",
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": base64.standard_b64encode(payload).decode("ascii"),
                    },
                }
            ],
            "pdf",
        )

    if mime in IMAGE_MIMES or suffix in {".jpg", ".jpeg", ".png", ".gif", ".webp"}:
        if len(data) > MAX_DOCUMENT_BYTES:
            raise Unreadable("image is too large to classify")
        media = mime if mime in IMAGE_MIMES else f"image/{suffix.lstrip('.').replace('jpg', 'jpeg')}"
        return (
            [
                {
                    "type": "image",
                    "source": {
                        "type": "base64",
                        "media_type": media,
                        "data": base64.standard_b64encode(data).decode("ascii"),
                    },
                }
            ],
            "image",
        )

    # A truncated upload or a mislabelled extension makes these parsers throw
    # library-specific exceptions. Every one of them is just "this file is not
    # readable" as far as the pipeline is concerned, and must not abort a batch
    # of a thousand other files.
    try:
        if mime in DOCX_MIMES or suffix == ".docx":
            text, modality = _docx_text(data), "docx"
        elif mime in XLSX_MIMES or suffix in {".xlsx", ".xlsm"}:
            text, modality = _xlsx_text(data), "xlsx"
        elif mime in CSV_MIMES or suffix in {".csv", ".tsv"}:
            text, modality = _csv_text(data), "csv"
        elif mime in TEXT_MIMES or suffix in {".txt", ".md", ".json", ".html"}:
            text, modality = data.decode("utf-8", errors="replace"), "text"
        else:
            raise Unreadable(f"unsupported type: {mime or suffix or 'unknown'}")
    except Unreadable:
        raise
    except Exception as exc:
        raise Unreadable(f"could not parse {suffix or mime} file: {exc}") from exc

    text = text.strip()
    if not text:
        raise Unreadable("no readable text in file")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n[... truncated for classification ...]"

    return ([{"type": "text", "text": f"File contents:\n\n{text}"}], modality)
