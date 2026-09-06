"""Offline rehearsal of a full scan / apply / undo cycle.

Runs the real pipeline, storage, naming, routing, journal and CLI rendering
against a temporary folder of badly named files, with a scripted classifier
standing in for the API. Nothing here calls Anthropic, so it costs nothing and
runs in CI.

    python scripts/demo.py
"""

from __future__ import annotations

import io
import shutil
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from rich.console import Console  # noqa: E402

from csfiler.classify import Classifier  # noqa: E402
from csfiler.cli import _render, _summary_line  # noqa: E402
from csfiler.config import load_config  # noqa: E402
from csfiler.journal import Journal  # noqa: E402
from csfiler.models import DocumentFacts  # noqa: E402
from csfiler.pipeline import Pipeline  # noqa: E402
from csfiler.storage import LocalStorage  # noqa: E402

console = Console()

# A realistic inbox: what actually lands in a Drive folder.
INBOX = [
    ("IMG_4471.pdf", "signed SOW"),
    ("scan (2).pdf", "mercury august statement"),
    ("Screenshot 2026-03-15 at 9.41.02 AM.png", "amex receipt"),
    ("Document(3).pdf", "1120-S for tax year 2025"),
    ("final proposal v2 FINAL.docx", "proposal for a prospect"),
    ("statement.pdf", "old entity insurance policy"),
    ("untitled folder export.pdf", "blurry unreadable scan"),
    ("standup-2026-08-12.txt", "meeting transcript"),
    ("2026-08-14_ThirdHorizon_SOW_signed.pdf", "already filed"),
]

# What the classifier would report for each, in order.
SCRIPT = {
    "IMG_4471.pdf": DocumentFacts(
        document_type="sow", document_date="2026-08-14", counterparty="Third Horizon",
        entity_named="Continuous Scale LLC", is_signed=True, confidence=0.96,
        summary="Signed statement of work for a Q3 engagement.",
        reasoning="Third Horizon letterhead, 'Statement of Work' title, signature block dated 8/14.",
    ),
    "scan (2).pdf": DocumentFacts(
        document_type="bank_statement", document_date="2026-08-31", period_year=2026,
        counterparty="Mercury", account_number="****4471", confidence=0.98,
        summary="Mercury operating account statement for August 2026.",
        reasoning="Mercury header, 'Statement Period Aug 1 - Aug 31, 2026', account ending 4471.",
    ),
    "Screenshot 2026-03-15 at 9.41.02 AM.png": DocumentFacts(
        document_type="card_receipt", document_date="2026-03-15", period_year=2026,
        counterparty="Amex", account_number="3782 822463 10005", confidence=0.91,
        summary="Card receipt for a software subscription charge.",
        reasoning="Amex logo, itemized total, card ending 0005.",
    ),
    "Document(3).pdf": DocumentFacts(
        document_type="tax_return", document_date="2026-03-15", period_year=2025,
        counterparty="IRS", entity_named="Continuous Scale LLC", confidence=0.97,
        summary="Form 1120-S filed for tax year 2025.",
        reasoning="Form 1120-S header, tax year box reads 2025, filed date stamped March 2026.",
    ),
    "final proposal v2 FINAL.docx": DocumentFacts(
        document_type="proposal", document_date="2026-09-01", counterparty="Acme Widgets",
        confidence=0.88, summary="Consulting proposal for a prospective client.",
        reasoning="Titled 'Proposal', scope and pricing sections, addressed to Acme Widgets.",
    ),
    "statement.pdf": DocumentFacts(
        document_type="insurance_policy", document_date="2026-01-01", period_year=2026,
        counterparty="Mercury", entity_named="Andover Consulting, LLC", confidence=0.93,
        summary="General liability declarations page.",
        reasoning="Declarations page listing the named insured and policy period.",
        concerns=["The named insured still reads Andover Consulting, LLC."],
    ),
    "untitled folder export.pdf": DocumentFacts(
        document_type="unknown", confidence=0.22,
        summary="An illegible scanned page.",
        reasoning="The page is a low-resolution scan; no header or totals are readable.",
    ),
    "standup-2026-08-12.txt": DocumentFacts(
        document_type="meeting_recording", document_date="2026-08-12", confidence=0.94,
        summary="Transcript of a daily standup.",
        reasoning="Timestamped speaker turns, no document structure.",
    ),
}


class ScriptedMessages:
    """Looks the answer up by filename, since scan order is not declaration order."""

    def __init__(self, script: dict[str, DocumentFacts]) -> None:
        self.script = script

    def parse(self, **kwargs):
        blocks = kwargs["messages"][0]["content"]
        # The filename hint is always the last block; for a text file the
        # contents block is also type "text", so take the last, not the first.
        hint = [b for b in blocks if b.get("type") == "text"][-1]["text"]
        name = hint.split("`")[1]
        if name not in self.script:
            raise AssertionError(f"demo has no scripted answer for {name}")
        return SimpleNamespace(
            parsed_output=self.script[name],
            stop_reason="end_turn",
            usage=SimpleNamespace(
                input_tokens=1800, output_tokens=180,
                cache_read_input_tokens=1400, cache_creation_input_tokens=0,
            ),
        )


class ScriptedClient:
    def __init__(self, script):
        self.messages = ScriptedMessages(script)


def make_pdf(text: str) -> bytes:
    """A genuinely valid one-page PDF, so the real extraction path runs."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    # Distinct metadata per document, so two different documents are not
    # byte-identical and correctly avoid the duplicate check.
    writer.add_metadata({"/Subject": text, "/Title": text})
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_png() -> bytes:
    """A 1x1 PNG - enough for the image branch to accept it."""
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmM"
        "IQAAAABJRU5ErkJggg=="
    )


def write_sample(path: Path, body: str) -> None:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        path.write_bytes(make_pdf(body))
    elif suffix == ".docx":
        path.write_bytes(make_docx(body))
    elif suffix == ".png":
        path.write_bytes(make_png())
    else:
        path.write_text(body + "\n" + "x" * 40, encoding="utf-8")


def tree(root: Path, prefix: str = "") -> list[str]:
    lines = []
    entries = sorted(root.iterdir(), key=lambda p: (p.is_file(), p.name))
    for i, entry in enumerate(entries):
        last = i == len(entries) - 1
        lines.append(f"{prefix}{'└── ' if last else '├── '}{entry.name}")
        if entry.is_dir():
            lines.extend(tree(entry, prefix + ("    " if last else "│   ")))
    return lines


def main() -> None:
    workdir = Path(tempfile.mkdtemp(prefix="csfiler-demo-"))
    inbox = workdir / "Inbox"
    inbox.mkdir()
    for name, body in INBOX:
        write_sample(inbox / name, body)

    config = load_config()
    storage = LocalStorage(workdir)
    classifier = Classifier(config, client=ScriptedClient(SCRIPT))
    journal = Journal(workdir / "journal.db")
    pipeline = Pipeline(config, storage, classifier, journal)

    console.rule("[bold]1. scan — reads and proposes, writes nothing")
    batch, proposals = pipeline.scan("Inbox")
    console.print()
    _render(proposals, show_all=True)
    console.print()
    _summary_line(proposals)
    console.print(f"\n[dim]{classifier.usage_summary()}[/dim]")

    console.print("\n[dim]Inbox is untouched after the scan:[/dim]")
    console.print("[dim]" + f"  {len(list(inbox.iterdir()))} files still in Inbox/" + "[/dim]")

    console.rule("\n[bold]2. apply — only the confident ones")
    auto = [p for p in proposals if p.disposition.value == "auto"]
    result = pipeline.apply(auto, batch)
    console.print(
        f"\n[green]{result.applied} filed[/green], {result.skipped} skipped, "
        f"[red]{result.failed} failed[/red]\n"
    )
    console.rule("\n[bold]Resulting BizBox tree")
    for line in tree(workdir):
        if "journal" not in line:
            console.print(f"[dim]{line}[/dim]")

    console.rule("\n[bold]3. undo — put everything back")
    undone = pipeline.undo(batch)
    console.print(f"\n[green]{undone.applied} reversed[/green]\n")
    remaining = sorted(p.name for p in inbox.iterdir())
    console.print(f"Inbox restored with {len(remaining)} files:")
    for name in remaining:
        console.print(f"  [dim]{name}[/dim]")

    journal.close()
    shutil.rmtree(workdir, ignore_errors=True)


if __name__ == "__main__":
    main()
